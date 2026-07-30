import os
import shutil
import subprocess
import pytest
from docx import Document
from ingestion.parsers.english import parse, BANK_ID
from ingestion.parsers.markers import MalformedItem
from ingestion.normalize import ResponseType

REAL_DOC_PATH = "Fraud support/FS Question Bank_English Proficiency V2.docx"

def test_english_parser_e2e():
    # 1. Parse the real document
    items = parse("Fraud support")
    
    # Assert counts
    assert len(items) == 105
    
    # Per-section counts
    sections = [item["section"] for item in items]
    assert sections.count("grammar") == 30
    assert sections.count("sentence_correction") == 30
    assert sections.count("macro") == 10
    assert sections.count("reading") == 25
    assert sections.count("closure") == 10
    
    # 2. Grammar golden anchor check
    grammar_q1 = next(item for item in items if item["section"] == "grammar" and item["source"]["number"] == 1)
    # The correct option should be 'b'
    correct_option = next(o for o in grammar_q1["options"] if o.is_correct)
    assert correct_option.letter == "b"
    assert grammar_q1["response_type"] == ResponseType.mcq_single
    
    # 3. Sentence correction check
    sc_items = [item for item in items if item["section"] == "sentence_correction"]
    for sc in sc_items:
        assert sc["stem"] != sc["model_answer"]
        assert sc["model_answer"]
        assert len(sc["options"]) == 0
        assert sc["response_type"] == ResponseType.open_text
        
    # 4. Macro checks
    macro_items = [item for item in items if item["section"] == "macro"]
    assert len(macro_items) == 10
    # Placeholder verification
    has_placeholder = False
    for m in macro_items:
        assert m["response_type"] == ResponseType.open_text
        assert m["model_answer"]
        if "{{ticket.requester.first_name}}" in m["stem"] and "{{ticket.requester.first_name}}" in m["model_answer"]:
            has_placeholder = True
    assert has_placeholder, "Macro placeholders like {{ticket.requester.first_name}} did not survive verbatim"
    
    # 5. Reading checks
    reading_items = [item for item in items if item["section"] == "reading"]
    assert len(reading_items) == 25
    
    # Check 5 passages x 5 items
    case_ids = [r["case_id"] for r in reading_items]
    unique_cases = sorted(list(set(case_ids)))
    assert len(unique_cases) == 5
    for case_id in unique_cases:
        case_qs = [r for r in reading_items if r["case_id"] == case_id]
        assert len(case_qs) == 5
        
    # Select-all verification (multiple correct options)
    has_multi = False
    total_correct_reading_checkmarks = 0
    for r in reading_items:
        correct_options = [o for o in r["options"] if o.is_correct]
        total_correct_reading_checkmarks += len(correct_options)
        if len(correct_options) > 1:
            assert r["response_type"] == ResponseType.mcq_multi
            has_multi = True
        else:
            assert r["response_type"] == ResponseType.mcq_single
    assert has_multi, "No select-all reading questions found"
    
    # 9. Marker census: exactly 38 ✅ across its 25 items
    assert total_correct_reading_checkmarks == 38
    
    # 6. Closure checks
    closure_items = [item for item in items if item["section"] == "closure"]
    assert len(closure_items) == 10
    for c in closure_items:
        assert c["response_type"] == ResponseType.hybrid
        assert len(c["options"]) == 3
        correct_opts = [o for o in c["options"] if o.is_correct]
        assert len(correct_opts) == 1
        assert c["model_answer"]
        assert c["case_id"].startswith("eng-closure-")
        assert c["case_title"]
        
    # 7. Grouping checks
    for item in items:
        assert item["bank"] == BANK_ID
        assert item["section"] is not None
        assert item["position"] is not None

def test_malformed_english_document(tmp_path):
    # Create a malformed docx in tmp_path and test direct parse raising MalformedItem
    malformed_path = tmp_path / "FS Question Bank_English Proficiency V2.docx"
    
    doc = Document(REAL_DOC_PATH)
    # Find the paragraph for Q1 in Part 2 and corrupt it (remove the Model Answer part)
    corrupted = False
    for p in doc.paragraphs:
        if p.text.startswith("Q1. Customer didn") and "Model Answer:" in p.text:
            p.text = "Q1. Customer didn’t sent the screenshot so we can’t verify nothing."
            corrupted = True
            break
            
    assert corrupted, "Could not find Q1 in Part 2 to corrupt"
    doc.save(malformed_path)
    
    with pytest.raises(MalformedItem) as exc:
        parse(str(tmp_path))
    assert "sentence_correction" in str(exc.value)
    assert "Question 1" in str(exc.value)

def test_cli_malformed_english_exits_nonzero():
    # Test that run.py CLI exits non-zero and prints the malformed error details when run on malformed data
    real_path = "Fraud support/FS Question Bank_English Proficiency V2.docx"
    backup_path = "Fraud support/FS Question Bank_English Proficiency V2.docx.bak"
    
    shutil.copy2(real_path, backup_path)
    try:
        # Load and corrupt Q1 in Part 2
        doc = Document(real_path)
        corrupted = False
        for p in doc.paragraphs:
            if p.text.startswith("Q1. Customer didn") and "Model Answer:" in p.text:
                p.text = "Q1. Customer didn’t sent the screenshot so we can’t verify nothing."
                corrupted = True
                break
        assert corrupted
        doc.save(real_path)
        
        # Run CLI in subprocess
        res = subprocess.run(
            ["python", "-m", "ingestion.run", "--banks", "english"],
            env=dict(os.environ, PYTHONPATH="."),
            capture_output=True,
            text=True
        )
        assert res.returncode != 0
        assert "sentence_correction" in res.stderr or "sentence_correction" in res.stdout
        assert "Question 1" in res.stderr or "Question 1" in res.stdout
    finally:
        if os.path.exists(backup_path):
            shutil.move(backup_path, real_path)
