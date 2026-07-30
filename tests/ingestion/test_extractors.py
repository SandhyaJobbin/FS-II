import pytest
from docx import Document
from ingestion.parsers.markers import (
    bold_correct_index, 
    trailing_check, 
    leading_check, 
    checkbox_pair, 
    split_lines, 
    MalformedItem, 
    clean_option_text
)
from ingestion.parsers.english import parse
from ingestion.manifest import SOURCE_DIR

def test_split_lines():
    text = "a) Option 1\nb) Option 2"
    assert split_lines(text) == ["a) Option 1", "b) Option 2"]
    assert split_lines("") == []

def test_clean_option_text():
    assert clean_option_text("a) Option A ✅") == ("Option A", "a)")
    assert clean_option_text("✅ A. Option 1") == ("Option 1", "A.")
    assert clean_option_text("☑ B. Option 2") == ("Option 2", "B.")
    assert clean_option_text("☐ C. Option 3") == ("Option 3", "C.")
    assert clean_option_text("Plain text option") == ("Plain text option", "")

def test_bold_correct_index(docx_bold_run, malformed_two_bold, malformed_no_marker):
    # Test valid bold run
    doc = Document(docx_bold_run)
    option_paras = doc.paragraphs[1:] # Skip question stem
    assert bold_correct_index(option_paras, "test.docx", "test", 1) == 1
    
    # Test malformed two bold
    doc_two = Document(malformed_two_bold)
    option_paras_two = doc_two.paragraphs[1:]
    with pytest.raises(MalformedItem) as exc_info:
        bold_correct_index(option_paras_two, "test.docx", "test", 2)
    assert "Expected exactly 1 bold correct option, found 2" in str(exc_info.value)
    
    # Test malformed no bold
    doc_none = Document(malformed_no_marker)
    option_paras_none = doc_none.paragraphs[1:]
    with pytest.raises(MalformedItem) as exc_info:
        bold_correct_index(option_paras_none, "test.docx", "test", 3)
    assert "Expected exactly 1 bold correct option, found 0" in str(exc_info.value)

def test_trailing_check_extractor(docx_trailing_check, malformed_no_marker):
    doc = Document(docx_trailing_check)
    option_lines = [p.text for p in doc.paragraphs[1:]]
    assert trailing_check(option_lines, "test.docx", "test", 1) == 1
    
    doc_none = Document(malformed_no_marker)
    option_lines_none = [p.text for p in doc_none.paragraphs[1:]]
    with pytest.raises(MalformedItem):
        trailing_check(option_lines_none, "test.docx", "test", 2)

def test_leading_check_extractor(docx_leading_check, malformed_no_marker):
    doc = Document(docx_leading_check)
    option_lines = [p.text for p in doc.paragraphs[1:]]
    assert leading_check(option_lines, "test.docx", "test", 1) == [0, 2]
    
    doc_none = Document(malformed_no_marker)
    option_lines_none = [p.text for p in doc_none.paragraphs[1:]]
    with pytest.raises(MalformedItem):
        leading_check(option_lines_none, "test.docx", "test", 2)

def test_checkbox_pair_extractor(docx_checkbox_pair, malformed_no_marker):
    doc = Document(docx_checkbox_pair)
    option_lines = [p.text for p in doc.paragraphs[1:]]
    assert checkbox_pair(option_lines, "test.docx", "test", 1) == [0]
    
    doc_none = Document(malformed_no_marker)
    option_lines_none = [p.text for p in doc_none.paragraphs[1:]]
    with pytest.raises(MalformedItem):
        checkbox_pair(option_lines_none, "test.docx", "test", 2)

def test_golden_anchor_english_grammar_q1():
    # Parse the real English proficiency file
    items = parse(SOURCE_DIR)
    # Reconcile expected size
    assert len(items) == 105
    
    # Q1 correct option is b (index 1)
    q1 = items[0]
    assert q1["stem"] == "The agent ___ the customer to provide additional evidence."
    assert len(q1["options"]) == 4
    
    # Option b is marked correct
    correct_opt = [o for o in q1["options"] if o.is_correct]
    assert len(correct_opt) == 1
    assert correct_opt[0].letter == "b"
    assert correct_opt[0].text == "asked"
