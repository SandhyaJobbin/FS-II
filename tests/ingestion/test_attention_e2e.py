import os
import shutil
import pytest
from docx import Document
from ingestion.parsers.attention import parse, BANK_ID, FILE_NAME
from ingestion.parsers.markers import MalformedItem
from ingestion.normalize import ResponseType

REAL_DOC_PATH = "Fraud support/FS Question Bank_Attention to Detail V2.docx"

def test_attention_parser_e2e():
    items = parse("Fraud support")
    
    # Assert counts
    assert len(items) == 160
    
    # Level divisions: 20 cases in L1, 20 cases in L2
    # 4 questions per case -> 80 questions per level
    l1_items = [item for item in items if item["level"] == "L1"]
    l2_items = [item for item in items if item["level"] == "L2"]
    assert len(l1_items) == 80
    assert len(l2_items) == 80
    
    # Cases count
    case_ids = {item["case_id"] for item in items}
    assert len(case_ids) == 40
    
    # Verify cases naming: attn-l1-case-01 to 20, attn-l2-case-01 to 20
    for case_num in range(1, 21):
        assert f"attn-l1-case-{case_num:02d}" in case_ids
        assert f"attn-l2-case-{case_num:02d}" in case_ids

    # Verify tables: 5 tables bound across cases
    unique_tables = []
    seen_cases = set()
    for item in items:
        if item["case_id"] not in seen_cases:
            seen_cases.add(item["case_id"])
            if item.get("tables"):
                unique_tables.extend(item["tables"])
    assert len(unique_tables) == 5
    
    for table in unique_tables:
        assert table["headers"]
        assert table["rows"]
        assert table["position"] in [1, 2, 3, 4, 5]
        
    # Verify verbatim tab names
    distinct_tabs = set()
    for item in items:
        if item.get("tabs"):
            for tab in item["tabs"]:
                distinct_tabs.add(tab["name"])
                
    assert len(distinct_tabs) >= 20
    assert "Review Information" in distinct_tabs
    assert "Account Profile" in distinct_tabs
    assert "Review Activity" in distinct_tabs
    
    # L2 Case 1 Q1 option splitting verification
    # Case 1 in L2 is attn-l2-case-01. Question 1
    l2_c1_q1 = next(item for item in items if item["case_id"] == "attn-l2-case-01" and item["source"]["number"] == 1)
    assert len(l2_c1_q1["options"]) == 4
    option_letters = [opt.letter for opt in l2_c1_q1["options"]]
    assert option_letters == ["A", "B", "C", "D"]
    
    # Single-select and Multi-select response types checks
    for item in items:
        assert item["bank"] == BANK_ID
        assert item["section"] in ["level_1", "level_2"]
        assert item["stem"]
        
        correct_count = sum(1 for opt in item["options"] if opt.is_correct)
        if item["response_type"] == ResponseType.mcq_single:
            assert correct_count == 1
        elif item["response_type"] == ResponseType.mcq_multi:
            assert correct_count >= 1

def test_malformed_attention_document_no_correct_options(tmp_path):
    # Corrupt a multi-select question to have all unchecked checkboxes (no ☑)
    malformed_path = tmp_path / FILE_NAME
    
    doc = Document(REAL_DOC_PATH)
    corrupted = False
    for p in doc.paragraphs:
        # Corrupt L2 Case 1 Q2 options (change ☑ to ☐)
        if "☑ A." in p.text and "Same device ID" in p.text:
            p.text = p.text.replace("☑ A.", "☐ A.").replace("☑ B.", "☐ B.").replace("☑ C.", "☐ C.")
            corrupted = True
            break
            
    assert corrupted, "Could not find multi-select option to corrupt"
    doc.save(malformed_path)
    
    with pytest.raises(MalformedItem) as exc:
        parse(str(tmp_path))
    assert "level_2" in str(exc.value)
    assert "Question 2" in str(exc.value)
    assert "No checked checkboxes" in str(exc.value)

def test_malformed_attention_document_multiple_single_correct(tmp_path):
    # Corrupt a single-select question to have two trailing ✅ marks
    malformed_path = tmp_path / FILE_NAME
    
    doc = Document(REAL_DOC_PATH)
    corrupted = False
    for p in doc.paragraphs:
        # Corrupt L1 Case 1 Q1 options (add ✅ to option A)
        if "A. The guest provided photos with the review" in p.text:
            p.text = "A. The guest provided photos with the review ✅\n" + \
                     "B. The review received a 3-star rating\n" + \
                     "C. The reviewer mentioned a full ocean view from a Standard Room ✅\n" + \
                     "D. The guest completed the booking"
            corrupted = True
            break
            
    assert corrupted, "Could not find single-select options to corrupt"
    doc.save(malformed_path)
    
    with pytest.raises(MalformedItem) as exc:
        parse(str(tmp_path))
    assert "level_1" in str(exc.value)
    assert "Question 1" in str(exc.value)
    assert "Expected exactly 1 trailing checkmark, found 2" in str(exc.value)
