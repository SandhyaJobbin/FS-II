import os
import pytest
from docx import Document

@pytest.fixture
def docx_trailing_check(tmp_path):
    path = os.path.join(tmp_path, "trailing_check.docx")
    doc = Document()
    doc.add_paragraph("Which of the following is correct?")
    doc.add_paragraph("a) Option 1")
    doc.add_paragraph("b) Option 2 ✅")
    doc.add_paragraph("c) Option 3")
    doc.save(path)
    return path

@pytest.fixture
def docx_leading_check(tmp_path):
    path = os.path.join(tmp_path, "leading_check.docx")
    doc = Document()
    doc.add_paragraph("Which of the following is correct?")
    doc.add_paragraph("✅ A. Option 1")
    doc.add_paragraph("B. Option 2")
    doc.add_paragraph("✅ C. Option 3")
    doc.save(path)
    return path

@pytest.fixture
def docx_checkbox_pair(tmp_path):
    path = os.path.join(tmp_path, "checkbox_pair.docx")
    doc = Document()
    doc.add_paragraph("Which of the following is correct?")
    doc.add_paragraph("☑ A. Correct Option")
    doc.add_paragraph("☐ B. Incorrect Option")
    doc.save(path)
    return path

@pytest.fixture
def docx_bold_run(tmp_path):
    path = os.path.join(tmp_path, "bold_run.docx")
    doc = Document()
    doc.add_paragraph("Which of the following is correct?")
    
    # Regular option
    p1 = doc.add_paragraph()
    p1.add_run("a) Regular Option")
    
    # Bold option
    p2 = doc.add_paragraph()
    r2 = p2.add_run("b) Bold Option")
    r2.bold = True
    
    # Regular option
    p3 = doc.add_paragraph()
    p3.add_run("c) Another Regular Option")
    
    doc.save(path)
    return path

@pytest.fixture
def malformed_missing_option(tmp_path):
    path = os.path.join(tmp_path, "missing_option.docx")
    doc = Document()
    doc.add_paragraph("Which of the following is correct?")
    # No options added
    doc.save(path)
    return path

@pytest.fixture
def malformed_no_marker(tmp_path):
    path = os.path.join(tmp_path, "no_marker.docx")
    doc = Document()
    doc.add_paragraph("Which of the following is correct?")
    doc.add_paragraph("a) Option 1")
    doc.add_paragraph("b) Option 2")
    doc.save(path)
    return path

@pytest.fixture
def malformed_two_bold(tmp_path):
    path = os.path.join(tmp_path, "two_bold.docx")
    doc = Document()
    doc.add_paragraph("Which of the following is correct?")
    
    p1 = doc.add_paragraph()
    r1 = p1.add_run("a) First Bold")
    r1.bold = True
    
    p2 = doc.add_paragraph()
    r2 = p2.add_run("b) Second Bold")
    r2.bold = True
    
    doc.save(path)
    return path
