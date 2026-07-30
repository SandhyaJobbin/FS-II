from typing import List, Dict, Any
from ingestion.normalize import Item, ResponseType
from ingestion.parsers.markers import MalformedItem

def validate_items(items: List[Item], manifest: Dict[str, Any], sections_run: List[str]):
    # Unique slugs tracking
    slug_to_item = {}
    
    # Track counts by section
    counts_by_section = {}
    
    # Check if tags file exists to determine if we should validate difficulty tier
    import os
    import json
    _tags_path = os.path.abspath(os.path.join(os.path.dirname(__file__), "..", "content", "difficulty-tags.json"))
    has_tags_file = os.path.exists(_tags_path)
    tags_keys = set()
    if has_tags_file:
        try:
            with open(_tags_path, "r", encoding="utf-8") as f:
                _data = json.load(f)
                tags_keys = {k for k in _data.keys() if not k.startswith("_")}
        except Exception:
            pass
            
    for item in items:
        # 1. Unique slugs check
        if item.id in slug_to_item:
            prev_item = slug_to_item[item.id]
            raise MalformedItem(
                item.source.file, 
                item.section, 
                item.source.number,
                f"Slug collision for ID '{item.id}' with question {prev_item.source.number} in section '{prev_item.section}'"
            )
        slug_to_item[item.id] = item
        
        # Difficulty tier validation when tags file is present
        CANONICAL_FILES = {
            "FS Question Bank_English Proficiency V2.docx",
            "FS Question Bank_Attention to Detail V2.docx",
            "FS Question Bank_Critical Thinking V2.docx"
        }
        if has_tags_file and item.source.file in CANONICAL_FILES:
            if item.id not in tags_keys:
                raise MalformedItem(
                    item.source.file, item.section, item.source.number,
                    f"Question ID '{item.id}' is missing from difficulty-tags.json"
                )
            if not item.difficulty_tier:
                raise MalformedItem(
                    item.source.file, item.section, item.source.number,
                    f"Question '{item.id}' has difficulty_tier set to None"
                )
            if item.difficulty_tier not in {"straightforward", "moderate", "complex"}:
                raise MalformedItem(
                    item.source.file, item.section, item.source.number,
                    f"Question '{item.id}' has invalid difficulty tier: {item.difficulty_tier}"
                )
        
        # Ingested section counts tracking
        counts_by_section[item.section] = counts_by_section.get(item.section, 0) + 1

        
        # 2. Options checks based on response type
        if item.response_type == ResponseType.mcq_single:
            if not item.options:
                raise MalformedItem(
                    item.source.file, item.section, item.source.number,
                    "MCQ single-select question has no options"
                )
            correct_count = sum(1 for o in item.options if o.is_correct)
            if correct_count != 1:
                raise MalformedItem(
                    item.source.file, item.section, item.source.number,
                    f"MCQ single-select question must have exactly 1 correct option, found {correct_count}"
                )
        elif item.response_type == ResponseType.mcq_multi:
            if not item.options:
                raise MalformedItem(
                    item.source.file, item.section, item.source.number,
                    "MCQ multi-select question has no options"
                )
            correct_count = sum(1 for o in item.options if o.is_correct)
            if correct_count < 1:
                raise MalformedItem(
                    item.source.file, item.section, item.source.number,
                    "MCQ multi-select question must have at least 1 correct option"
                )
                
    # 3. Check section and level counts against manifest expectations
    for bank_name, bank_meta in manifest.items():
        if "sections" in bank_meta:
            for sect, expected_count in bank_meta["sections"].items():
                if sect in sections_run:
                    actual_count = counts_by_section.get(sect, 0)
                    if actual_count != expected_count:
                        raise ValueError(
                            f"Count mismatch in section '{sect}' of bank '{bank_name}': "
                            f"Expected {expected_count} questions, but found {actual_count}"
                        )
        if "levels" in bank_meta:
            for lvl, expected_cases in bank_meta["levels"].items():
                sect_name = f"level_{lvl[1].lower()}" # level_1 or level_2
                if sect_name in sections_run:
                    actual_count = counts_by_section.get(sect_name, 0)
                    expected_questions = expected_cases * 4
                    if actual_count != expected_questions:
                        raise ValueError(
                            f"Count mismatch in section '{sect_name}' of bank '{bank_name}': "
                            f"Expected {expected_questions} questions, but found {actual_count}"
                        )
                        
    # 4. Marker census reconciliation
    import os
    from docx import Document
    for bank_name, bank_meta in manifest.items():
        run_this_bank = any(item.bank == bank_name for item in items)
        if run_this_bank and "markers" in bank_meta:
            file_path = os.path.join("Fraud support", bank_meta["file"])
            if os.path.exists(file_path):
                doc = Document(file_path)
                expected_markers = bank_meta["markers"]
                if bank_name == "attention":
                    c_yes = sum(p.text.count('✅') for p in doc.paragraphs)
                    c_chk = sum(p.text.count('☑') for p in doc.paragraphs)
                    c_unchk = sum(p.text.count('☐') for p in doc.paragraphs)
                    
                    if c_yes != expected_markers.get("trailing_check"):
                        raise ValueError(f"Marker census mismatch for 'trailing_check' (✅) in bank '{bank_name}': expected {expected_markers.get('trailing_check')}, found {c_yes}")
                    if c_chk != expected_markers.get("checkbox_checked"):
                        raise ValueError(f"Marker census mismatch for 'checkbox_checked' (☑) in bank '{bank_name}': expected {expected_markers.get('checkbox_checked')}, found {c_chk}")
                    if c_unchk != expected_markers.get("checkbox_unchecked"):
                        raise ValueError(f"Marker census mismatch for 'checkbox_unchecked' (☐) in bank '{bank_name}': expected {expected_markers.get('checkbox_unchecked')}, found {c_unchk}")
                elif bank_name == "critical":
                    c_yes = sum(p.text.count('✅') for p in doc.paragraphs)
                    if c_yes != expected_markers.get("leading_check"):
                        raise ValueError(f"Marker census mismatch for 'leading_check' (✅) in bank '{bank_name}': expected {expected_markers.get('leading_check')}, found {c_yes}")
                        
    return True
