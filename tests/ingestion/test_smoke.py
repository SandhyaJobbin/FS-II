import os
from docx import Document

def test_docx_trailing_check(docx_trailing_check):
    assert os.path.exists(docx_trailing_check)
    doc = Document(docx_trailing_check)
    texts = [p.text for p in doc.paragraphs]
    assert len(texts) == 4
    assert any(text.endswith("✅") for text in texts)

def test_docx_leading_check(docx_leading_check):
    assert os.path.exists(docx_leading_check)
    doc = Document(docx_leading_check)
    texts = [p.text for p in doc.paragraphs]
    assert len(texts) == 4
    assert texts[1].startswith("✅")
    assert texts[3].startswith("✅")

def test_docx_checkbox_pair(docx_checkbox_pair):
    assert os.path.exists(docx_checkbox_pair)
    doc = Document(docx_checkbox_pair)
    texts = [p.text for p in doc.paragraphs]
    assert len(texts) == 3
    assert any(text.startswith("☑") for text in texts)
    assert any(text.startswith("☐") for text in texts)

def test_docx_bold_run(docx_bold_run):
    assert os.path.exists(docx_bold_run)
    doc = Document(docx_bold_run)
    bold_runs = []
    for p in doc.paragraphs:
        for r in p.runs:
            if r.bold:
                bold_runs.append(r.text)
    assert len(bold_runs) == 1
    assert "Bold Option" in bold_runs[0]

def test_malformed_missing_option(malformed_missing_option):
    assert os.path.exists(malformed_missing_option)
    doc = Document(malformed_missing_option)
    assert len(doc.paragraphs) == 1

def test_malformed_no_marker(malformed_no_marker):
    assert os.path.exists(malformed_no_marker)
    doc = Document(malformed_no_marker)
    texts = [p.text for p in doc.paragraphs]
    assert len(texts) == 3
    assert not any("✅" in text or "☑" in text or "☐" in text for text in texts)

def test_malformed_two_bold(malformed_two_bold):
    assert os.path.exists(malformed_two_bold)
    doc = Document(malformed_two_bold)
    bold_runs = []
    for p in doc.paragraphs:
        for r in p.runs:
            if r.bold:
                bold_runs.append(r.text)
    assert len(bold_runs) == 2
